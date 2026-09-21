import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser(description="Regenerate the plain-language three-week plan.")
parser.add_argument(
    "--source",
    type=Path,
    default=REPOSITORY_ROOT / "docs/deliverables/plans/Ke_hoach_3_tuan_Bach_Hoa_Sim_Tim_phan_cong_lai.docx",
)
parser.add_argument(
    "--output",
    type=Path,
    default=REPOSITORY_ROOT / "docs/deliverables/plans/Ke_hoach_3_tuan_Bach_Hoa_Sim_Tim_khong_ma_cong_viec.docx",
)
args = parser.parse_args()
SOURCE = args.source.resolve()
OUTPUT = args.output.resolve()


def set_font(run, size=9, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_text(paragraph, text, size=10.5, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)


def set_cell_text(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, cm):
    cell.width = Cm(cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.first_child_found_in("w:tcW")
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(Cm(cm).emu / 635)))
    tcW.set(qn("w:type"), "dxa")


def style_table(table, body_size, widths=None, center_cols=()):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if widths and c_idx < len(widths):
                set_width(cell, widths[c_idx])
            align = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or c_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            for p in cell.paragraphs:
                p.alignment = align
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, size=body_size + (0.35 if r_idx == 0 else 0), bold=True if r_idx == 0 else None)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def delete_table_column(table, index):
    for tr in table._tbl.tr_lst:
        cells = tr.tc_lst
        if index < len(cells):
            tr.remove(cells[index])
    grid = table._tbl.tblGrid
    cols = grid.gridCol_lst
    if index < len(cols):
        grid.remove(cols[index])


def disable_repeated_header(row):
    trPr = row._tr.get_or_add_trPr()
    for node in list(trPr):
        if node.tag == qn("w:tblHeader"):
            trPr.remove(node)


def fill_table(table, rows, size, center_cols=()):
    assert len(table.rows) == len(rows), (len(table.rows), len(rows))
    for r_idx, row_data in enumerate(rows):
        assert len(table.rows[r_idx].cells) == len(row_data), (r_idx, len(table.rows[r_idx].cells), len(row_data))
        for c_idx, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or c_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(table.rows[r_idx].cells[c_idx], value, size=size + (0.35 if r_idx == 0 else 0), bold=(r_idx == 0), align=align)


doc = Document(SOURCE)
assert len(doc.tables) == 8

# Diễn đạt phần hướng dẫn bằng tiếng Việt dễ hiểu, không dùng mã đầu việc.
paragraph_updates = {
    11: "Kế hoạch này kế thừa phần nền tảng đã có và tập trung hoàn thiện sản phẩm trong ba tuần. Mỗi công việc đều có người phụ trách chính, người rà soát, điều kiện hoàn thành và minh chứng cụ thể.",
    13: "Nguyên tắc phân công: mỗi thành viên trực tiếp lập trình một phần chức năng có thể trình diễn; người phụ trách chịu trách nhiệm từ triển khai đến kiểm thử và minh chứng; người rà soát kiểm tra trước khi gộp vào bản chung.",
    15: "Minh chứng dùng để đánh giá gồm: đầu việc đã hoàn thành, lịch sử thay đổi mã nguồn, yêu cầu gộp mã, kết quả kiểm thử, hình ảnh hoặc video chạy chức năng và tài liệu bàn giao.",
    17: "Nhóm ưu tiên hoàn thành trọn vẹn luồng nghiệp vụ cốt lõi trước, sau đó mới làm các phần mở rộng. Công việc chưa đáp ứng điều kiện hoàn thành sẽ không được tính là đã xong.",
    20: "Mức Must là bắt buộc để nghiệm thu; Should là nên hoàn thành nếu không ảnh hưởng phần bắt buộc; Could chỉ thực hiện khi còn thời gian; Won’t là không triển khai trong ba tuần này.",
    26: "Nguồn tham khảo: tài liệu quản lý dự án của Bách Hóa Sim Tím trên GitHub, truy cập ngày 21/09/2026. Đường dẫn: https://github.com/zomboXx/bach-hoa-sim-tim/tree/main/docs/project/governance",
}
for idx, text in paragraph_updates.items():
    set_paragraph_text(doc.paragraphs[idx], text)
set_paragraph_text(doc.paragraphs[14], "1.1 Mục tiêu phân bổ đóng góp kỹ thuật", size=14, bold=True)
set_paragraph_text(doc.paragraphs[21], "3.1 Ma trận MoSCoW theo đầu việc và người phụ trách", size=14, bold=True)
set_paragraph_text(doc.paragraphs[24], "4.1 Điều kiện chuyển tuần", size=14, bold=True)

# Bảng 1: phân công trách nhiệm.
assignment_rows = [
    ["STT", "Thành viên", "Vai trò", "Phạm vi sở hữu", "Công việc chịu trách nhiệm", "Người rà soát chính"],
    ["1", "Nguyễn Đức Phát\n24110296", "Nhóm trưởng\nPhụ trách kỹ thuật và bảo mật máy chủ", "Đăng nhập và phân quyền phía máy chủ; xử lý thanh toán, hóa đơn và xuất lô theo hạn dùng; lưu kết quả đào tạo; kiến trúc và tích hợp", "Trực tiếp lập trình đăng nhập, phiên làm việc, chống giả mạo yêu cầu và phân quyền; xử lý thanh toán, hóa đơn và xuất lô theo hạn dùng; xây dựng dịch vụ lưu phiên và kết quả đào tạo; viết kiểm thử cho phần mình. Đồng thời điều phối việc gộp mã, tích hợp và hồ sơ bàn giao.", "Nguyễn Văn Trung"],
    ["2", "Nguyễn Văn Trung\n24110365", "Cơ sở dữ liệu và máy chủ", "Cơ sở dữ liệu, danh mục, nhận hàng, tồn kho, đồng bộ xung đột và duyệt điều chỉnh", "Lập trình cấu trúc dữ liệu, cập nhật cơ sở dữ liệu, dữ liệu mẫu, danh mục, nhận hàng, tồn kho, lô và hạn dùng; xử lý đồng bộ xung đột và duyệt điều chỉnh; viết kiểm thử giao dịch, phục hồi dữ liệu và các dịch vụ mình phụ trách.", "Nguyễn Đức Phát"],
    ["3", "Nguyễn Văn Thi\n24110334", "Giao diện bán hàng", "Kết nối giao diện với máy chủ; đăng nhập; bán hàng; hóa đơn; khuyến mãi; báo cáo và giao diện hiện trường", "Lập trình lớp kết nối dữ liệu; hoàn thiện màn hình đăng nhập, bán hàng, hóa đơn, khuyến mãi và báo cáo; bổ sung trạng thái tải, lỗi và dữ liệu rỗng; tối ưu giao diện trình diễn trên máy tính và điện thoại.", "Lê Văn Chiến"],
    ["4", "Lê Văn Chiến\n24110170", "Di động, mất mạng, đào tạo và kiểm thử tự động", "Kiểm kê trên điện thoại; hàng đợi khi mất mạng; kịch bản đào tạo; kiểm thử tự động và kiểm thử tổng thể", "Lập trình giao diện kiểm kê phù hợp điện thoại, hàng đợi khi mất mạng và xử lý kết nối lại; hoàn thiện một kịch bản đào tạo; xây dựng kiểm thử tự động cho các luồng chính và tổng hợp báo cáo lỗi.", "Nguyễn Văn Thi"],
]
fill_table(doc.tables[0], assignment_rows, 8.7, center_cols=(0, 5))
style_table(doc.tables[0], 8.7, widths=[0.65, 2.15, 2.35, 3.35, 7.1, 2.0], center_cols=(0, 5))

# Bảng 2: tỷ trọng đóng góp kỹ thuật.
contribution_rows = [
    ["Thành viên", "Phần lập trình sở hữu", "Tỷ trọng kỹ thuật dự kiến", "Minh chứng bắt buộc"],
    ["Nguyễn Đức Phát", "Đăng nhập, phân quyền và bảo mật phía máy chủ; thanh toán, hóa đơn và xuất lô theo hạn dùng; lưu kết quả đào tạo; tích hợp hệ thống", "25%", "Lịch sử thay đổi mã nguồn, yêu cầu gộp mã, kiểm thử quyền và giao dịch, video chạy chức năng"],
    ["Nguyễn Văn Trung", "Cơ sở dữ liệu; danh mục; nhận hàng; tồn kho; đồng bộ xung đột; duyệt điều chỉnh", "25%", "Cấu trúc dữ liệu, dữ liệu mẫu, dịch vụ máy chủ, kiểm thử giao dịch và phục hồi dữ liệu"],
    ["Nguyễn Văn Thi", "Kết nối giao diện; đăng nhập; bán hàng; hóa đơn; khuyến mãi; báo cáo; giao diện trình diễn", "25%", "Màn hình hoàn chỉnh, kết nối dữ liệu thật, kiểm thử giao diện, hình ảnh hoặc video trình diễn"],
    ["Lê Văn Chiến", "Kiểm kê trên điện thoại; làm việc khi mất mạng; kịch bản đào tạo; kiểm thử tự động", "25%", "Mã nguồn chức năng, bộ kiểm thử tự động, báo cáo kết quả và video chạy trên điện thoại"],
]
fill_table(doc.tables[1], contribution_rows, 8.9, center_cols=(2,))
style_table(doc.tables[1], 8.9, widths=[2.7, 8.1, 2.1, 5.3], center_cols=(2,))

# Bảng 3: kế hoạch ba tuần, chỉ dùng tên công việc.
schedule_rows = [
    ["Tuần", "Thời gian", "Người phụ trách", "Công việc", "Kết quả và điều kiện hoàn thành", "Người rà soát", "Xác nhận GVHD"],
    ["1", "21–27/09/2026", "Nguyễn Đức Phát", "Chốt phạm vi và hoàn thiện đăng nhập phân quyền\n- Điền phân công thật và chốt phạm vi.\n- Lập trình đăng nhập, phiên làm việc, chống giả mạo yêu cầu và phân quyền phía máy chủ.\n- Viết kiểm thử quyền; thống nhất cách kết nối giao diện, máy chủ và phần đào tạo.", "Danh sách công việc có người phụ trách và người rà soát; kiểm tra tự động hoạt động; người dùng sai quyền bị từ chối; kiểm thử bảo mật đạt yêu cầu.", "Nguyễn Văn Trung"],
    ["1", "21–27/09/2026", "Nguyễn Văn Trung", "Hoàn thiện cơ sở dữ liệu và chức năng danh mục\n- Chốt mô hình dữ liệu và từ điển dữ liệu.\n- Tạo cập nhật cấu trúc dữ liệu, dữ liệu mẫu, kiểm tra tình trạng hệ thống.\n- Lập trình danh mục hàng hóa, nhà cung cấp và tìm kiếm theo mã vạch.", "Có thể tạo cơ sở dữ liệu mới từ đầu; dữ liệu mẫu hợp lệ; chức năng thêm, xem, sửa, xóa và tìm kiếm chạy được; có kiểm thử.", "Nguyễn Đức Phát"],
    ["1", "21–27/09/2026", "Nguyễn Văn Thi", "Kết nối giao diện với máy chủ và hoàn thiện đăng nhập\n- Tạo lớp kết nối dữ liệu dùng chung.\n- Nối màn hình đăng nhập với máy chủ.\n- Bổ sung trạng thái đang tải, lỗi và dữ liệu rỗng.", "Đăng nhập bằng dữ liệu thật; thông báo lỗi rõ ràng; giao diện không còn phụ thuộc dữ liệu giả trong luồng nghiệm thu.", "Lê Văn Chiến"],
    ["1", "21–27/09/2026", "Lê Văn Chiến", "Lập kế hoạch kiểm thử và ma trận truy vết\n- Viết các kịch bản kiểm thử cho đăng nhập, danh mục, nhận hàng và bán hàng.\n- Thiết lập bộ kiểm thử tự động ban đầu.\n- Ghi lỗi theo mẫu thống nhất.", "Mỗi yêu cầu bắt buộc có ít nhất một kịch bản kiểm thử; bộ kiểm thử có thể chạy lặp lại; lỗi có người xử lý và hạn sửa.", "Nguyễn Văn Thi"],
    ["2", "28/09–04/10/2026", "Nguyễn Văn Trung", "Hoàn thiện nhận hàng, tồn kho, lô hàng và hạn dùng\n- Lập trình phiếu nhận hàng và các dòng hàng.\n- Cập nhật tồn kho, lô, hạn dùng và lịch sử biến động trong cùng một giao dịch.\n- Viết kiểm thử trường hợp thành công và hoàn tác khi lỗi.", "Nhận hàng làm tăng tồn đúng; sai dữ liệu không để lại trạng thái dở dang; có lịch sử biến động và kiểm thử giao dịch.", "Nguyễn Đức Phát"],
    ["2", "28/09–04/10/2026", "Nguyễn Văn Thi", "Hoàn thiện giao diện bán hàng, hóa đơn, khuyến mãi và báo cáo\n- Nối giỏ hàng với dữ liệu thật.\n- Hiển thị hóa đơn và kết quả thanh toán.\n- Bổ sung khuyến mãi cơ bản và báo cáo tổng hợp.", "Giao diện gọi đúng dịch vụ bán hàng phía máy chủ; trạng thái thành công và thất bại rõ ràng; báo cáo phản ánh đúng dữ liệu mẫu.", "Lê Văn Chiến"],
    ["2", "28/09–04/10/2026", "Nguyễn Đức Phát", "Hoàn thiện xử lý thanh toán, hóa đơn và xuất kho phía máy chủ\n- Tạo hóa đơn và các dòng hóa đơn trong một giao dịch.\n- Chọn lô sắp hết hạn trước và giảm tồn kho.\n- Hoàn tác toàn bộ khi có lỗi; viết kiểm thử tích hợp.", "Bán hàng tạo đúng hóa đơn, trừ đúng lô và tồn; lỗi không làm mất dữ liệu; kiểm thử thành công, thất bại và quyền truy cập đều đạt.", "Nguyễn Văn Trung"],
    ["2", "28/09–04/10/2026", "Lê Văn Chiến", "Tự động hóa kiểm thử luồng nhận hàng và bán hàng\n- Kiểm thử toàn bộ luồng từ nhận hàng đến bán hàng.\n- Kiểm tra quy tắc xuất lô sắp hết hạn trước và hoàn tác khi lỗi.\n- Xuất báo cáo kết quả kiểm thử.", "Các luồng chính chạy tự động ổn định; lỗi có bằng chứng; báo cáo nêu rõ số ca đạt, không đạt và nguyên nhân.", "Nguyễn Văn Thi"],
    ["3", "05–11/10/2026", "Lê Văn Chiến", "Hoàn thiện kiểm kê trên điện thoại, làm việc khi mất mạng và kịch bản đào tạo\n- Tối ưu màn hình kiểm kê cho điện thoại.\n- Lưu tạm thao tác khi mất mạng và gửi lại khi có kết nối.\n- Hoàn thiện một kịch bản đào tạo bằng Godot.", "Kiểm kê dùng tốt trên điện thoại; thao tác không mất khi ngắt mạng; dữ liệu gửi lại đúng; kịch bản đào tạo có thể trình diễn từ đầu đến cuối.", "Nguyễn Văn Thi"],
    ["3", "05–11/10/2026", "Nguyễn Văn Trung", "Hoàn thiện đồng bộ xung đột và duyệt điều chỉnh tồn\n- Phát hiện dữ liệu thay đổi đồng thời.\n- Ngăn gửi trùng thao tác.\n- Xây dựng quy trình tạo, duyệt và từ chối điều chỉnh tồn; kiểm tra phục hồi dữ liệu.", "Xung đột được thông báo rõ; gửi lại không tạo bản ghi trùng; chỉ người có quyền mới duyệt; dữ liệu có thể sao lưu và phục hồi.", "Nguyễn Đức Phát"],
    ["3", "05–11/10/2026", "Nguyễn Văn Thi", "Hoàn thiện giao diện hiện trường và bản trình diễn\n- Tinh chỉnh màn hình kiểm kê, bán hàng và báo cáo trên điện thoại.\n- Hiển thị trạng thái mất mạng, chờ gửi và xung đột.\n- Chuẩn bị dữ liệu và kịch bản trình diễn.", "Giao diện dễ thao tác; trạng thái hệ thống dễ hiểu; bản trình diễn chạy liên tục theo kịch bản nghiệm thu.", "Lê Văn Chiến"],
    ["3", "05–11/10/2026", "Nguyễn Đức Phát", "Lưu kết quả đào tạo, kiểm thử tổng thể và hoàn thiện bản bàn giao\n- Lập trình lưu phiên và kết quả đào tạo.\n- Rà soát các thay đổi, sửa lỗi tích hợp và kiểm thử lại.\n- Hoàn thiện hướng dẫn cài đặt, vận hành, minh chứng và bản chạy thử cuối.", "Kết quả đào tạo lưu đúng người dùng; toàn bộ phần bắt buộc đạt kiểm thử; tài liệu đủ để cài đặt, trình diễn, đối chiếu đóng góp và vấn đáp.", "Nguyễn Văn Trung"],
]
date_short = {
    "21–27/09/2026": "21–27/09",
    "28/09–04/10/2026": "28/09–04/10",
    "05–11/10/2026": "05–11/10",
}
for row in schedule_rows[1:]:
    row[1] = date_short[row[1]]
    row.append("")
fill_table(doc.tables[2], schedule_rows, 8.05, center_cols=(0, 1, 2, 5, 6))
style_table(doc.tables[2], 8.05, widths=[0.6, 1.45, 1.9, 5.8, 5.45, 1.75, 1.25], center_cols=(0, 1, 2, 5, 6))
disable_repeated_header(doc.tables[2].rows[0])

# Bảng 4: MoSCoW tóm tắt.
moscow_summary_rows = [
    ["Mức ưu tiên", "Nội dung", "Cách xử lý", "Người điều phối"],
    ["Must\nBắt buộc", "Đăng nhập và phân quyền; cơ sở dữ liệu; danh mục; nhận hàng; tồn kho; bán hàng; hóa đơn; xuất lô theo hạn dùng; kiểm thử toàn bộ luồng; tài liệu và bản chạy thử cuối.", "Bảo vệ phạm vi. Mỗi phần việc có người phụ trách, người rà soát và điều kiện hoàn thành.", "Nguyễn Đức Phát"],
    ["Should\nNên có", "Khuyến mãi cơ bản; kiểm kê trên điện thoại; làm việc khi mất mạng; đồng bộ xung đột; duyệt điều chỉnh tồn; một kịch bản đào tạo; lưu kết quả đào tạo.", "Chỉ triển khai sau khi các luồng bắt buộc đã ổn định. Có kiểm thử và minh chứng như phần bắt buộc.", "Người phụ trách từng phần"],
    ["Could\nCó thể có", "Bộ lọc báo cáo nâng cao; tinh chỉnh giao diện; mở rộng kiểm thử tự động; thêm hình ảnh và âm thanh cho phần đào tạo.", "Thực hiện khi còn thời gian; không được làm chậm phần bắt buộc và phần nên có.", "Thành viên còn thời gian"],
    ["Won’t\nChưa làm", "Bán hàng khi mất mạng trên nhiều thiết bị; ứng dụng cài riêng và thiết bị bán hàng chuyên dụng; mô hình dịch vụ cho nhiều cửa hàng; thành viên và đổi điểm; toàn bộ các chương đào tạo còn lại.", "Ghi rõ ngoài phạm vi ba tuần; không nhận thêm nếu chưa có quyết định đổi phạm vi của cả nhóm.", "Nguyễn Đức Phát"],
]
fill_table(doc.tables[3], moscow_summary_rows, 8.9)
style_table(doc.tables[3], 8.9, widths=[2.0, 8.4, 5.0, 2.6])

# Bảng 5: MoSCoW chi tiết. Xóa hẳn cột mã công việc.
detailed = doc.tables[4]
delete_table_column(detailed, 1)
moscow_detail_rows = [
    ["Ưu tiên", "Công việc", "Người phụ trách", "Người rà soát", "Tuần", "Điều kiện hoàn thành"],
    ["Must", "Quy trình, phân công, quản lý đầu việc, yêu cầu gộp mã và kiểm tra tự động", "Nguyễn Đức Phát", "Cả nhóm", "1", "Phân công rõ; mọi thay đổi được rà soát; kiểm tra tự động chạy ổn định."],
    ["Must", "Chốt phạm vi, mô hình dữ liệu và từ điển dữ liệu", "Nguyễn Văn Trung", "Nguyễn Đức Phát", "1", "Phạm vi được thống nhất; dữ liệu và quan hệ được mô tả đầy đủ."],
    ["Must", "Máy chủ, cập nhật cấu trúc dữ liệu, danh mục và nhà cung cấp", "Nguyễn Văn Trung", "Nguyễn Đức Phát", "1", "Có thể khởi tạo dữ liệu từ đầu; chức năng danh mục chạy được và có kiểm thử."],
    ["Must", "Đăng nhập, phiên làm việc, chống giả mạo yêu cầu và phân quyền phía máy chủ", "Nguyễn Đức Phát", "Nguyễn Văn Trung", "1", "Đăng nhập đúng; sai quyền bị từ chối; kiểm thử bảo mật đạt."],
    ["Must", "Kết nối giao diện với máy chủ và đăng nhập", "Nguyễn Văn Thi", "Lê Văn Chiến", "1", "Giao diện dùng dữ liệu thật, có trạng thái tải, lỗi và dữ liệu rỗng."],
    ["Must", "Kiểm thử nền tảng và ma trận truy vết", "Lê Văn Chiến", "Nguyễn Văn Thi", "1", "Mỗi yêu cầu bắt buộc có kịch bản kiểm thử và bằng chứng kết quả."],
    ["Must", "Nhận hàng, tồn kho, lô, hạn dùng và biến động", "Nguyễn Văn Trung", "Nguyễn Đức Phát", "2", "Tồn kho cập nhật đúng; dữ liệu lỗi được hoàn tác; có lịch sử biến động."],
    ["Must", "Xử lý thanh toán, hóa đơn, xuất lô theo hạn dùng và giảm tồn", "Nguyễn Đức Phát", "Nguyễn Văn Trung", "2", "Hóa đơn và tồn kho nhất quán; chọn đúng lô; có kiểm thử thành công và thất bại."],
    ["Must", "Giao diện bán hàng, hóa đơn và báo cáo", "Nguyễn Văn Thi", "Lê Văn Chiến", "2", "Hoàn thành luồng bán hàng bằng dữ liệu thật; báo cáo phản ánh đúng kết quả."],
    ["Must", "Kiểm thử toàn bộ luồng nhận hàng và bán hàng", "Lê Văn Chiến", "Nguyễn Văn Thi", "2", "Các luồng chính chạy tự động; báo cáo nêu rõ số ca đạt và không đạt."],
    ["Must", "Kiểm thử lại, phân quyền và phục hồi dữ liệu", "Lê Văn Chiến", "Nguyễn Đức Phát", "3", "Không còn lỗi nghiêm trọng; quyền truy cập đúng; phục hồi dữ liệu thành công."],
    ["Must", "Báo cáo, liên kết minh chứng, bản bàn giao và chuẩn bị vấn đáp", "Nguyễn Đức Phát", "Cả nhóm", "3", "Tài liệu đủ để cài đặt, trình diễn, đối chiếu đóng góp và trả lời vấn đáp."],
    ["Should", "Khuyến mãi cơ bản", "Nguyễn Văn Thi", "Nguyễn Đức Phát", "2", "Áp dụng đúng điều kiện; tổng tiền chính xác; có kiểm thử."],
    ["Should", "Kiểm kê trên điện thoại và hàng đợi khi mất mạng", "Lê Văn Chiến", "Nguyễn Văn Thi", "3", "Thao tác dùng được trên điện thoại; dữ liệu không mất khi ngắt kết nối."],
    ["Should", "Đồng bộ xung đột và duyệt điều chỉnh tồn", "Nguyễn Văn Trung", "Nguyễn Đức Phát", "3", "Không tạo bản ghi trùng; xung đột rõ ràng; chỉ đúng quyền mới được duyệt."],
    ["Should", "Một kịch bản đào tạo bằng Godot", "Lê Văn Chiến", "Nguyễn Văn Thi", "3", "Kịch bản chạy từ đầu đến cuối và có video minh chứng."],
    ["Should", "Dịch vụ lưu phiên và kết quả đào tạo tách biệt", "Nguyễn Đức Phát", "Nguyễn Văn Trung", "3", "Kết quả lưu đúng người dùng; có kiểm thử và tài liệu kết nối."],
    ["Could", "Bộ lọc báo cáo và tinh chỉnh giao diện ngoài tiêu chí nghiệm thu", "Nguyễn Văn Thi", "Lê Văn Chiến", "3", "Chỉ thực hiện sau khi phần bắt buộc và phần nên có đã ổn định."],
    ["Could", "Mở rộng kiểm thử tự động và hình ảnh đào tạo ngoài kịch bản cam kết", "Lê Văn Chiến", "Nguyễn Văn Thi", "3", "Không làm chậm việc sửa lỗi và hoàn thiện hồ sơ."],
    ["Won’t", "Bán hàng khi mất mạng trên nhiều thiết bị; ứng dụng cài riêng và thiết bị bán hàng; mô hình dịch vụ cho nhiều cửa hàng; thành viên và đổi điểm; toàn bộ các chương đào tạo còn lại", "Không phân công", "Nguyễn Đức Phát", "-", "Ghi rõ ngoài phạm vi ba tuần; không triển khai trong đợt này."],
]
fill_table(detailed, moscow_detail_rows, 8.35, center_cols=(0, 2, 3, 4))
style_table(detailed, 8.35, widths=[1.1, 6.15, 2.35, 2.35, 0.75, 5.5], center_cols=(0, 2, 3, 4))

# Bảng 6: nhịp làm việc.
cadence_rows = [
    ["Hoạt động", "Tần suất", "Người chủ trì", "Kết quả bắt buộc"],
    ["Lập kế hoạch tuần", "Đầu mỗi tuần", "Nguyễn Đức Phát", "Chốt mục tiêu, người phụ trách, người rà soát và hạn hoàn thành."],
    ["Cập nhật tiến độ", "Hằng ngày, 15 phút", "Luân phiên", "Nêu việc đã làm, việc sẽ làm và vướng mắc cần hỗ trợ."],
    ["Rà soát thay đổi mã nguồn", "Trong ngày", "Người rà soát được phân công", "Góp ý rõ ràng; kiểm tra tự động đạt; thay đổi đủ điều kiện gộp."],
    ["Trình diễn nội bộ", "Cuối mỗi tuần", "Người phụ trách từng phần", "Chạy chức năng thật; lưu hình ảnh hoặc video và kết quả kiểm thử."],
    ["Rút kinh nghiệm", "Sau trình diễn", "Nguyễn Đức Phát", "Chốt điều cần giữ, điều cần sửa và hành động cho tuần kế tiếp."],
]
fill_table(doc.tables[5], cadence_rows, 9.0, center_cols=(1, 2))
style_table(doc.tables[5], 9.0, widths=[3.5, 2.6, 3.35, 8.25], center_cols=(1, 2))

# Bảng 7: cổng kiểm soát theo tuần.
gate_rows = [
    ["Mốc kiểm soát", "Điều kiện đạt và minh chứng", "Quyết định nếu chưa đạt"],
    ["Cuối tuần 1", "Phạm vi, người phụ trách, người rà soát và nền tảng kỹ thuật đã rõ; đăng nhập và danh mục dùng dữ liệu thật. Minh chứng: danh sách công việc, mô hình dữ liệu, kiểm thử, hình ảnh hoặc video.", "Dừng phần mở rộng để sửa nền tảng trước khi chuyển sang luồng nghiệp vụ."],
    ["Cuối tuần 2", "Nhận hàng, tồn kho và bán hàng chạy liên tục; hóa đơn và giảm tồn đúng; kiểm thử tự động đạt. Minh chứng: video, báo cáo kiểm thử, lịch sử thay đổi và yêu cầu gộp mã.", "Ưu tiên sửa luồng bắt buộc trước khi chuyển sang làm ổn định và bàn giao."],
    ["Cuối tuần 3", "Không còn lỗi nghiêm trọng; phần bắt buộc hoàn thành; bản chạy thử và tài liệu bàn giao đầy đủ. Minh chứng: biên bản kiểm thử, hướng dẫn cài đặt, ma trận truy vết và video.", "Ghi rõ phần chưa hoàn thành; chỉ nộp khi đã thống nhất phạm vi thực tế và chuẩn bị được phần vấn đáp."],
]
fill_table(doc.tables[6], gate_rows, 9.0)
style_table(doc.tables[6], 9.0, widths=[2.4, 10.3, 5.3])

# Đồng bộ kiểu chữ cho bảng chữ ký, không thay đổi nội dung.
style_table(doc.tables[7], 10.0, widths=[9.0, 9.0], center_cols=(0, 1))

doc.core_properties.title = "Kế hoạch thực hiện ba tuần – Bách Hóa Sim Tím"
doc.core_properties.subject = "Phân công công việc và ưu tiên MoSCoW, trình bày bằng tên công việc không dùng mã"
doc.core_properties.comments = "Bản chỉnh sửa ngày 21/09/2026: bỏ mã đầu việc, diễn đạt rõ trách nhiệm lập trình của từng thành viên."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
